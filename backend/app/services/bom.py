from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import xlrd


class BomError(Exception):
    pass


@dataclass
class ParsedLine:
    line_no: int
    designators: str = ""
    quantity_per_board: int = 0
    quantity_required: int = 0
    lcsc_code: str = ""
    mpn: str = ""
    brand: str = ""
    footprint: str = ""
    value: str = ""
    params: str = ""
    product_name: str = ""
    category: str = ""


@dataclass
class ParsedBom:
    name: str
    format: str
    purchase_sets: int = 1
    order_no: str = ""
    lines: list[ParsedLine] = field(default_factory=list)


EDA_COLUMNS = {
    "no.",
    "quantity",
    "footprint",
    "lcsc price",
    "manufacturer",
    "manufacturer part",
    "value",
}

VALUE_IN_PARAMS = re.compile(r"(?:阻值|容值|电感量|电感值)[:：]\s*(\S+)")
VALUE_TEXT = re.compile(
    r"(ω|ω|Ω|ohm|[pnuµmμ]f|[pnuµmμ]h|\d\s*k$|\d\s*m$)",
    re.IGNORECASE,
)


def parse_upload(filename: str, data: bytes) -> ParsedBom:
    suffix = Path(filename or "").suffix.lower()
    try:
        if suffix == ".xlsx":
            matrix = _read_xlsx(data)
        elif suffix == ".xls":
            matrix = _read_xls(data)
        else:
            raise BomError("只支持 .xls 或 .xlsx")
    except BomError:
        raise
    except Exception as exc:
        raise BomError("文件无法读取") from exc
    if not matrix:
        raise BomError("无法识别文件格式")
    return _parse_matrix(matrix, filename)


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _read_xlsx(data: bytes) -> list[list[str]]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheet = workbook.active
    row_count = sheet.max_row or 0
    col_count = sheet.max_column or 0
    matrix = [
        [_cell(sheet.cell(row, col).value) for col in range(1, col_count + 1)]
        for row in range(1, row_count + 1)
    ]
    for merged in sheet.merged_cells.ranges:
        value = matrix[merged.min_row - 1][merged.min_col - 1]
        for row in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                matrix[row - 1][col - 1] = value
    return matrix


def _read_xls(data: bytes) -> list[list[str]]:
    workbook = xlrd.open_workbook(file_contents=data)
    sheet = workbook.sheet_by_index(0)
    matrix = [
        [_cell(sheet.cell_value(row, col)) for col in range(sheet.ncols)]
        for row in range(sheet.nrows)
    ]
    for row_start, row_end, col_start, col_end in getattr(sheet, "merged_cells", []) or []:
        value = matrix[row_start][col_start]
        for row in range(row_start, row_end):
            for col in range(col_start, col_end):
                matrix[row][col] = value
    return matrix


def _header_key(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip()).lower()


def _header_index(row: list[str]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, cell in enumerate(row):
        key = _header_key(cell)
        if key and key not in found:
            found[key] = index
    return found


def _find_header(matrix: list[list[str]], required: set[str]) -> int | None:
    for index, row in enumerate(matrix[:40]):
        keys = {_header_key(cell) for cell in row if cell}
        if required <= keys:
            return index
    return None


def _pick(row: list[str], index: dict[str, int], *names: str) -> str:
    for name in names:
        position = index.get(_header_key(name))
        if position is not None and position < len(row):
            return row[position].strip()
    return ""


def _parse_int(value: str, default: int = 0) -> int:
    text = (value or "").strip().replace(",", "")
    if not text:
        return default
    try:
        return int(float(text))
    except ValueError:
        return default


def _looks_like_value(value: str) -> bool:
    text = value.strip().lower().replace("μ", "u").replace("µ", "u")
    if not text or not re.search(r"\d", text):
        return False
    return VALUE_TEXT.search(text) is not None


def _parse_matrix(matrix: list[list[str]], filename: str) -> ParsedBom:
    head = "\n".join(" ".join(row) for row in matrix[:15])
    quote_header = _find_header(matrix, {"商品编号", "designator"})
    eda_header = _find_header(matrix, {"manufacturer part", "lcsc price", "value"})
    if "立创商城BOM报价单" in head or quote_header is not None:
        if quote_header is None:
            raise BomError("无法识别文件格式")
        return _parse_quote(matrix, quote_header, filename)
    if eda_header is not None:
        return _parse_eda(matrix, eda_header, filename)
    raise BomError("无法识别文件格式")


def _meta(matrix: list[list[str]], header_index: int) -> tuple[str, int, str]:
    name = ""
    purchase_sets = 1
    order_no = ""
    for row in matrix[:header_index]:
        for cell in row:
            name_match = re.search(r"BOM名称[:：]\s*(.+)", cell)
            if name_match:
                name = name_match.group(1).strip()
            sets_match = re.search(r"采购套数[:：]\s*(\d+)", cell)
            if sets_match:
                purchase_sets = int(sets_match.group(1))
            order_match = re.search(r"BOM配单编号[:：]\s*(\S+)", cell)
            if order_match:
                order_no = order_match.group(1).strip()
    return name, purchase_sets, order_no


def _parse_quote(matrix: list[list[str]], header_index: int, filename: str) -> ParsedBom:
    name, purchase_sets, order_no = _meta(matrix, header_index)
    columns = _header_index(matrix[header_index])
    lines: list[ParsedLine] = []
    for row in matrix[header_index + 1 :]:
        if not any(row):
            continue
        per_board = _parse_int(_pick(row, columns, "Quantity"))
        required = per_board * purchase_sets
        params = "；".join(
            part
            for part in (_pick(row, columns, "原参数"), _pick(row, columns, "参数"))
            if part
        )
        value_match = VALUE_IN_PARAMS.search(params)
        name_cell = _pick(row, columns, "Name")
        if value_match:
            value = value_match.group(1)
        elif _looks_like_value(name_cell):
            value = name_cell
        else:
            value = ""
        mall_footprint = _pick(row, columns, "封装")
        original_footprint = _pick(row, columns, "Footprint")
        footprint = mall_footprint or original_footprint
        if mall_footprint and original_footprint and mall_footprint != original_footprint:
            extra = f"原始封装: {original_footprint}"
            params = f"{params}；{extra}" if params else extra
        line = ParsedLine(
            line_no=_parse_int(_pick(row, columns, "ID", "No."), default=len(lines) + 1),
            designators=_pick(row, columns, "Designator"),
            quantity_per_board=per_board,
            quantity_required=required,
            lcsc_code=_pick(row, columns, "商品编号"),
            mpn=_pick(row, columns, "型号") or _pick(row, columns, "Manufacturer Part"),
            brand=_pick(row, columns, "品牌") or _pick(row, columns, "Manufacturer"),
            footprint=footprint,
            value=value,
            params=params,
            product_name=_pick(row, columns, "商品名称"),
            category=_pick(row, columns, "目录"),
        )
        if any((line.lcsc_code, line.mpn, line.designators, line.product_name, line.value)):
            lines.append(line)
    if not lines:
        raise BomError("没有可导入的物料明细")
    return ParsedBom(
        name=name or Path(filename).stem,
        format="lcsc_quote",
        purchase_sets=purchase_sets,
        order_no=order_no,
        lines=lines,
    )


def _parse_eda(matrix: list[list[str]], header_index: int, filename: str) -> ParsedBom:
    header = matrix[header_index]
    columns = _header_index(header)
    lines: list[ParsedLine] = []
    for row in matrix[header_index + 1 :]:
        if not any(row):
            continue
        mpn = _pick(row, columns, "Manufacturer Part")
        value = _pick(row, columns, "Value")
        footprint = _pick(row, columns, "Footprint")
        brand = _pick(row, columns, "Manufacturer")
        if not any((mpn, value, footprint, brand)):
            continue
        extras: list[str] = []
        for index, title in enumerate(header):
            key = _header_key(title)
            if not key or key in EDA_COLUMNS:
                continue
            cell = row[index].strip() if index < len(row) else ""
            if cell:
                extras.append(f"{title.strip()}: {cell}")
        quantity = _parse_int(_pick(row, columns, "Quantity"))
        lines.append(
            ParsedLine(
                line_no=_parse_int(_pick(row, columns, "No."), default=len(lines) + 1),
                quantity_per_board=quantity,
                quantity_required=quantity,
                mpn=mpn,
                brand=brand,
                footprint=footprint,
                value=value,
                params="；".join(extras),
            )
        )
    if not lines:
        raise BomError("没有可导入的物料明细")
    return ParsedBom(
        name=Path(filename).stem,
        format="easyeda_v131",
        purchase_sets=1,
        lines=lines,
    )
